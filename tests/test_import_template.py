"""The downloadable sample workbook: well-formed, documented, and round-trips through the importer."""
from io import BytesIO

from openpyxl import load_workbook

from conftest import login


def test_requires_cs(client, user):
    login(client, 'bob@test.com')
    assert client.get('/cs/import/template.xlsx').status_code in (302, 403)


def test_template_structure(client, cs_user):
    login(client, 'cs@test.com')
    r = client.get('/cs/import/template.xlsx')
    assert r.status_code == 200
    assert 'spreadsheetml' in r.headers['Content-Type']
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames == ['Travellers', 'How to fill']
    ws = wb['Travellers']
    headers = [c.value for c in ws[1]]
    assert headers[:6] == ['poster_name', 'traveler_name', 'on_behalf_of', 'role', 'origin', 'destination']
    assert ws.max_row == 4                                  # header + 3 example rows
    guide = wb['How to fill']
    assert any('seeking_help' in str(c.value) for row in guide.iter_rows() for c in row)


def test_template_round_trips_through_the_importer(client, cs_user):
    login(client, 'cs@test.com')
    data = client.get('/cs/import/template.xlsx').data
    r = client.post('/cs/import', data={'file': (BytesIO(data), 'template.xlsx'), 'default_source': 'excel'},
                    content_type='multipart/form-data')
    html = r.data.decode()
    assert r.status_code == 200
    assert 'Priya Patel' in html and 'Ravi Kumar' in html and 'Anita D' in html
    assert 'QR573' in html                                  # flight parsed
