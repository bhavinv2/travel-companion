"""Sales CRM, second part: the WhatsApp CRM switch and its contract, Instagram / Facebook,
leads that sit too long, alerts, the per-person numbers, the inbox check, and bringing in the
past (site enquiries and the team's sales workbook)."""
import hashlib
import hmac
import io
import json
import time
from datetime import date, datetime, timedelta
from types import SimpleNamespace

import openpyxl
import pytest

from conftest import login, make_user
from app import db as _db
from app.models import (CrmChannelCheck, CrmIdentity, CrmLead, CrmPerson, CrmPipeline, CrmSource, CrmStage,
                        CrmTask, CrmTouchpoint, Notification)
from app.services import crm, crm_import, crm_reports, meta, wacrm


def stage(pipe_key, stage_key):
    return (CrmStage.query.join(CrmPipeline).filter(CrmPipeline.key == pipe_key, CrmStage.key == stage_key)
            .one())


def ti_lead(phone='+919876543210', **kw):
    return crm.capture('ti-quote', name=kw.pop('name', 'Ramesh Kumar'), phone=phone, utm={}, **kw)


@pytest.fixture
def wa_keys(monkeypatch):
    monkeypatch.setenv('WACRM_URL', 'https://crm.example/crm')
    monkeypatch.setenv('WACRM_API_KEY', 'wacrm_live_test')


# ============================================================================================ whatsapp crm

def test_keys_alone_send_nothing(app, wa_keys, monkeypatch):
    calls = []
    monkeypatch.setattr(wacrm, '_call', lambda *a, **k: calls.append(a) or {})
    lead = ti_lead()
    assert crm.wa_mode() == 'off' and calls == []
    assert crm.send_whatsapp(lead, 'Hello')['mode'] == 'link'
    assert calls == []


def test_contacts_mode_adds_contacts_but_sends_nothing(app, wa_keys, monkeypatch):
    crm.save_settings(whatsapp_mode='contacts')
    calls = []

    def fake(method, path, payload=None, params=None):
        calls.append((method, path, payload))
        return {'id': 'c-1', 'tags': []} if method == 'POST' else {}
    monkeypatch.setattr(wacrm, '_call', fake)
    lead = ti_lead()
    assert lead.person.wacrm_contact_id == 'c-1'
    assert calls[0][:2] == ('POST', '/api/v1/contacts') and 'tags' not in calls[0][2]
    # the product's tag, and where the lead is ("TI: Cold lead") for broadcasts and automations
    assert calls[1] == ('PATCH', '/api/v1/contacts/c-1', {'tags': ['TI: Cold lead', 'ti']})
    assert crm.send_whatsapp(lead, 'Hello')['mode'] == 'link'


def test_contact_tags_are_merged_not_replaced(app, wa_keys, monkeypatch):
    calls = []

    def fake(method, path, payload=None, params=None):
        calls.append((method, path, payload))
        return {'id': 'c-9', 'tags': [{'name': 'vip'}, {'name': 'ti'}]}
    monkeypatch.setattr(wacrm, '_call', fake)
    assert wacrm.upsert_contact('+919876543210', tags=['ti']) == 'c-9'
    assert len(calls) == 1                                  # nothing missing: no PATCH at all
    wacrm.upsert_contact('+919876543210', tags=['ti', 'diwali'])
    assert calls[-1] == ('PATCH', '/api/v1/contacts/c-9', {'tags': ['diwali', 'ti', 'vip']})


def test_template_params_are_strings_and_language_optional(app, wa_keys, monkeypatch):
    sent = []
    monkeypatch.setattr(wacrm, '_call', lambda m, p, payload=None, params=None: sent.append(payload) or {})
    wacrm.send_template('+919876543210', 'intro', params=['Ramesh', 3])
    assert sent[0]['template'] == {'name': 'intro', 'params': ['Ramesh', '3']}


def test_html_answer_means_wrong_address(app, wa_keys, monkeypatch):
    class R:
        status_code = 404

        def json(self):
            raise ValueError
    monkeypatch.setattr(wacrm.requests, 'request', lambda *a, **k: R())
    with pytest.raises(wacrm.WacrmError, match='WACRM_URL'):
        wacrm.me()


def test_admin_switches_and_tests_the_connection(app, client, admin_user, wa_keys, monkeypatch):
    monkeypatch.setattr(wacrm, 'me', lambda: {'account': {'name': 'Admin'},
                                              'key': {'scopes': ['contacts:read', 'contacts:write', 'messages:send']}})
    login(client, admin_user.email)
    r = client.post('/admin/crm/setup/whatsapp', data={'action': 'test'}, follow_redirects=True)
    assert b'Connected to the WhatsApp CRM account' in r.data
    client.post('/admin/crm/setup/whatsapp', data={'action': 'mode', 'mode': 'send'})
    assert crm.wa_mode() == 'send'
    page = client.get('/admin/crm/setup?tab=connect').get_data(as_text=True)
    assert 'Connect replies' in page


def test_registering_the_webhook_keeps_its_secret(app, client, admin_user, wa_keys, monkeypatch):
    monkeypatch.setattr(wacrm, 'register_webhook', lambda url: {'id': 'w1', 'secret': 'whsec_abc', 'url': url})
    login(client, admin_user.email)
    # the test server is not public: refused with a reason
    r = client.post('/admin/crm/setup/whatsapp', data={'action': 'webhook'}, follow_redirects=True)
    assert b'Do this on the live site' in r.data
    app.config['PREFERRED_URL_SCHEME'] = 'https'
    monkeypatch.setattr('app.routes.crm._public_url', lambda ep: 'https://example.org/api/crm/wacrm-webhook')
    client.post('/admin/crm/setup/whatsapp', data={'action': 'webhook'})
    assert wacrm.webhook_secret() == 'whsec_abc'


def test_failed_delivery_is_flagged_to_the_sender(app, client, cs_user, monkeypatch):
    monkeypatch.setenv('WACRM_WEBHOOK_SECRET', 'whsec_x')
    lead = ti_lead()
    crm.log_touch(lead, 'whatsapp', 'out', body='Hi', user=cs_user, status='sent', external_id='wamid.7')
    _db.session.commit()
    body = json.dumps({'event': 'message.status_updated',
                       'data': {'whatsapp_message_id': 'wamid.7', 'status': 'failed'}}).encode()
    t = int(time.time())
    sig = hmac.new(b'whsec_x', str(t).encode() + b'.' + body, hashlib.sha256).hexdigest()
    client.post('/api/crm/wacrm-webhook', data=body, headers={'X-Wacrm-Signature': 't=%d,v1=%s' % (t, sig),
                                                               'Content-Type': 'application/json'})
    _db.session.expire_all()
    assert CrmTouchpoint.query.filter_by(external_id='wamid.7').one().status == 'failed'
    assert Notification.query.filter_by(user_id=cs_user.id, type='crm_failed').count() == 1


# ============================================================================================ identities

def test_handles_are_one_person_however_written(app):
    a = crm.capture('instagram', name='Priya', identities=[('instagram', '@Priya.R')], utm={})
    b = crm.capture('instagram', identities=[('instagram', 'https://www.instagram.com/priya.r/?hl=en')],
                    message='again', utm={})
    assert a.id == b.id and CrmPerson.query.count() == 1
    assert a.person.identity('instagram') == 'priya.r'


def test_quick_add_with_only_an_instagram_handle(app, client, cs_user):
    login(client, cs_user.email)
    crm.ensure_defaults()
    r = client.post('/cs/crm/add', data={'product': 'TI', 'name': '', 'instagram': '@meena.j',
                                         'source': 'instagram', 'owner': 'me'})
    assert r.status_code == 302
    lead = CrmLead.query.one()
    assert lead.person.display_name == 'meena.j' and lead.created_by_id == cs_user.id
    assert lead.pipeline.key == 'ti-standard'


# ============================================================================================ meta

def meta_post(client, payload, secret='appsecret'):
    raw = json.dumps(payload).encode()
    sig = 'sha256=' + hmac.new(secret.encode(), raw, hashlib.sha256).hexdigest()
    return client.post('/api/crm/meta-webhook', data=raw,
                       headers={'X-Hub-Signature-256': sig, 'Content-Type': 'application/json'})


@pytest.fixture
def meta_keys(monkeypatch):
    monkeypatch.setenv('META_APP_SECRET', 'appsecret')
    monkeypatch.setenv('META_PAGE_TOKEN', 'pagetoken')
    monkeypatch.setenv('META_VERIFY_TOKEN', 'verify-me')
    monkeypatch.setattr(meta, 'profile', lambda sid, ch: {'name': 'Anita Rao', 'username': 'anita.rao'})


def test_meta_parses_messages_and_comments():
    payload = {'object': 'instagram', 'entry': [{'id': 'PAGE', 'messaging': [
        {'sender': {'id': 'IG1'}, 'recipient': {'id': 'PAGE'}, 'message': {'mid': 'm1', 'text': 'Price for my mom?'}},
        {'sender': {'id': 'PAGE'}, 'message': {'mid': 'm2', 'text': 'our echo', 'is_echo': True}}],
        'changes': [{'field': 'comments', 'value': {'id': 'c1', 'text': 'Interested!', 'from': {'id': 'IG2', 'username': 'raj'}}}]}]}
    evs = meta.parse_events(payload)
    assert [(e['kind'], e['sender_id']) for e in evs] == [('message', 'IG1'), ('comment', 'IG2')]
    fb = {'object': 'page', 'entry': [{'id': 'PAGE', 'changes': [{'field': 'feed', 'value': {
        'item': 'comment', 'verb': 'add', 'comment_id': 'fc1', 'message': 'How much?', 'from': {'id': 'F1', 'name': 'Ravi'}}}]}]}
    assert meta.parse_events(fb)[0]['channel'] == 'facebook'


def test_meta_webhook_verifies_and_files_messages(app, client, meta_keys):
    r = client.get('/api/crm/meta-webhook?hub.mode=subscribe&hub.verify_token=verify-me&hub.challenge=42')
    assert r.data == b'42'
    assert client.get('/api/crm/meta-webhook?hub.mode=subscribe&hub.verify_token=nope&hub.challenge=1').status_code == 403
    msg = lambda mid, text: {'object': 'instagram', 'entry': [{'id': 'PAGE', 'messaging': [  # noqa: E731
        {'sender': {'id': 'IG1'}, 'message': {'mid': mid, 'text': text}}]}]}
    assert meta_post(client, msg('m1', 'Hi, is my mother covered for diabetes?')).get_json()['stored'] == 1
    lead = CrmLead.query.one()
    assert lead.source.key == 'instagram' and lead.person.name == 'Anita Rao'
    assert lead.person.identity('ig_id') == 'IG1' and lead.person.identity('instagram') == 'anita.rao'
    meta_post(client, msg('m2', 'Also my father'))
    meta_post(client, msg('m2', 'Also my father'))                    # Meta resends: not twice
    _db.session.expire_all()
    assert CrmLead.query.count() == 1
    assert CrmTouchpoint.query.filter_by(direction='in', channel='instagram').count() == 2
    assert meta_post(client, msg('m3', 'x'), secret='wrong').status_code == 401


def test_reply_on_instagram_from_the_lead(app, client, cs_user, meta_keys, monkeypatch):
    sent = []
    monkeypatch.setattr(meta, 'send_text', lambda rid, text: sent.append((rid, text)) or 'mid.9')
    meta_post(client, {'object': 'instagram', 'entry': [{'id': 'PAGE', 'messaging': [
        {'sender': {'id': 'IG1'}, 'message': {'mid': 'm1', 'text': 'hello'}}]}]})
    lead = CrmLead.query.one()
    login(client, cs_user.email)
    page = client.get('/cs/crm/lead/%d' % lead.id).get_data(as_text=True)
    assert 'data-via="instagram"' in page
    d = client.post('/cs/crm/lead/%d/send' % lead.id, json={'text': 'Hi Anita!', 'via': 'instagram'}).get_json()
    assert d['mode'] == 'sent' and sent == [('IG1', 'Hi Anita!')]
    assert CrmTouchpoint.query.filter_by(direction='out', channel='instagram').one().external_id == 'mid.9'


def test_instagram_reply_refused_without_a_conversation(app, client, cs_user, meta_keys):
    lead = ti_lead()
    login(client, cs_user.email)
    r = client.post('/cs/crm/lead/%d/send' % lead.id, json={'text': 'Hi', 'via': 'instagram'})
    assert r.status_code == 400 and 'reply only' in r.get_json()['error']


# ============================================================================================ alerts

def test_new_lead_rings_the_owners_bell(app, cs_user, admin_user):
    lead = ti_lead()
    n = Notification.query.filter_by(user_id=cs_user.id, type='crm_lead').one()
    assert 'Ramesh Kumar' in n.title and n.link.endswith('/cs/crm/lead/%d' % lead.id)
    crm.capture('brand-new-page', phone='+14155550123', utm={})
    assert Notification.query.filter_by(user_id=admin_user.id, type='crm_source').count() == 1


def test_lead_waiting_too_long_is_flagged_once(app, cs_user):
    lead = ti_lead()                                          # cold lead: 24 hours allowed
    assert lead.stuck_hours is None
    lead.entered_stage_at = datetime.utcnow() - timedelta(hours=30)
    _db.session.commit()
    assert lead.stuck_hours == 6
    assert crm.run_sla_alerts()['crm_stuck_alerts'] == 1
    assert crm.run_sla_alerts()['crm_stuck_alerts'] == 0      # once per stage
    assert Notification.query.filter_by(user_id=cs_user.id, type='crm_stuck').count() == 1
    crm.move_stage(lead, stage('ti-standard', 'prospect'))
    assert lead.sla_alerted_at is None and lead.stuck_hours is None


def test_stuck_leads_show_on_the_todo_list(app, client, cs_user):
    lead = ti_lead()
    lead.entered_stage_at = datetime.utcnow() - timedelta(hours=50)
    _db.session.commit()
    login(client, cs_user.email)
    assert 'Leads waiting too long' in client.get('/cs/crm/today').get_data(as_text=True)
    assert 'waiting' in client.get('/cs/crm/?owner=all').get_data(as_text=True)


def test_expiry_reminder_is_due_before_the_policy_ends(app):
    lead = ti_lead(ends_on=date.today() + timedelta(days=40))
    step = SimpleNamespace(trigger='before_end_days', trigger_value=10)
    due = crm._due(step, lead, datetime.utcnow())
    assert (due + crm.IST).date() == date.today() + timedelta(days=30) and (due + crm.IST).hour == 9
    customer = stage('ti-standard', 'customer')
    assert [s.code for s in customer.steps][-1] == 'ER'
    assert '{ends}' not in crm.fill('Ends on {ends}', lead)


# ============================================================================================ numbers

def test_per_person_numbers(app, client, cs_user, admin_user):
    lead = ti_lead()
    login(client, cs_user.email)
    client.post('/cs/crm/lead/%d/send' % lead.id, json={'text': 'Hello', 'task_id': lead.open_task.id})
    client.post('/cs/crm/lead/%d/call' % lead.id, json={'outcome': 'connected'})
    totals, daily = crm_reports.kpis(datetime.utcnow() - timedelta(days=1))
    t = totals[cs_user.id]
    assert (t['leads'], t['messages'], t['calls'], t['talked'], t['done']) == (1, 1, 1, 1, 1)
    assert t['codes']['TI'] == 1 and t['first_reply_h'] is not None
    assert client.get('/cs/crm/insights').status_code == 302          # the dashboard is the admin's
    client.post('/auth/logout')
    login(client, admin_user.email)
    page = client.get('/admin/crm/insights?who=%d' % cs_user.id).get_data(as_text=True)
    assert 'Csagent&#39;s numbers' in page and 'Day by day' in page and 'Messages by kind' in page
    csv = client.get('/admin/crm/insights.csv?days=7')
    assert csv.headers['Content-Type'].startswith('text/csv') and b'Messages sent' in csv.data


# ============================================================================================ checker

@pytest.fixture
def ten_fifteen(monkeypatch):
    fixed = datetime(2026, 10, 12, 10, 15)                  # 10:15 in India -> slot 20
    monkeypatch.setattr(crm_reports, 'ist_now', lambda: fixed)
    return fixed


def test_inbox_check(app, client, cs_user, ten_fifteen):
    login(client, cs_user.email)
    page = client.get('/cs/crm/checks').get_data(as_text=True)
    assert 'Now: 10:00' in page and 'IND WhatsApp' in page
    assert client.post('/cs/crm/checks/tick', json={'all': True}).get_json()['all']
    assert CrmChannelCheck.query.filter_by(slot=20).count() == len(crm.crm_settings()['check_channels'])
    assert crm_reports.current_slot_done(crm.crm_settings()['check_channels'])
    assert client.post('/cs/crm/checks/tick', json={'slot': 20, 'channel': 'Tawk'}).get_json()['ticked'] is False
    assert client.post('/cs/crm/checks/tick', json={'slot': 21, 'channel': 'Tawk'}).status_code == 400   # the future
    assert client.post('/cs/crm/checks/tick', json={'slot': 15, 'channel': 'Tawk'}).get_json()['ticked'] is True
    cov = crm_reports.coverage(ten_fifteen.date(), crm.crm_settings()['check_channels'])
    assert cov['due'] == 6 and cov['done'] == 0                # 07:00-09:30 not fully checked


def test_admin_edits_the_inbox_list(app, client, admin_user):
    login(client, admin_user.email)
    client.post('/admin/crm/setup/whatsapp', data={'action': 'channels', 'channels': 'Tawk\nIND WhatsApp\n'})
    assert crm.crm_settings()['check_channels'] == ['Tawk', 'IND WhatsApp']


# ============================================================================================ the past

def test_backfill_brings_in_past_enquiries_once(app, user):
    from app.models import CompanionRequest, ContactMessage, InsuranceQuote, SahayakBooking
    today = date.today()
    _db.session.add_all([
        InsuranceQuote(email='a@example.com', phone='+919000000001', insurance_type='visitors', citizenship='IND',
                       start_date=today + timedelta(days=20), end_date=today + timedelta(days=80), travellers=[{'age': '66'}]),
        InsuranceQuote(email='b@example.com', phone='+919000000002', insurance_type='visitors', citizenship='IND',
                       start_date=today - timedelta(days=20), end_date=today + timedelta(days=40), travellers=[],
                       created_at=datetime.utcnow() - timedelta(days=50)),
        ContactMessage(name='Old', email='c@example.com', phone='+919000000003', message='Need help', topic='general',
                       created_at=datetime.utcnow() - timedelta(days=60)),
        SahayakBooking(service_key='phlebotomy', service_name='Blood test', patient_name='Amma', phone='+919000000004',
                       address='Delhi', status='completed'),
    ])
    trip = CompanionRequest(user_id=user.id, travel_type='air', trip_type='one_way', source='organic',
                            flying_from='Hyderabad (HYD)', destination='Dallas (DFW)', from_date=today + timedelta(days=30))
    trip.set_status('open')
    _db.session.add(trip)
    _db.session.commit()
    out = crm_import.backfill()
    assert (out['insurance quotes'], out['contact messages'], out['sahayak bookings'], out['trip posts']) == (2, 1, 1, 1)
    quotes = {l.person.phone: l for l in CrmLead.query.filter_by(ref_type='insurance_quotes')}
    live, gone = quotes['+919000000001'], quotes['+919000000002']
    assert live.status == 'open' and live.open_task is not None and live.ends_on
    assert gone.status == 'lost' and gone.fallout_reason == 'Trip date passed' and gone.open_task is None
    old = CrmLead.query.filter_by(ref_type='contact_messages').one()
    assert old.temperature == 'EL' and old.open_task is None
    assert CrmLead.query.filter_by(ref_type='sahayak_bookings').one().status == 'won'
    assert Notification.query.filter_by(type='crm_lead').count() == 0          # quietly
    again = crm_import.backfill()
    assert again['already'] == 5 and CrmLead.query.count() == 5


def workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Cold Leads'
    ws.append([None, 1, 'HYD DFW NRI TC', '1.804.497.6837'])
    ws.append([None, 2, 'BLR DTW NRI TC', 'https://www.instagram.com/some.one?igsh=x'])
    ws.append([None, 3, 'nobody here', 'call later'])
    p = wb.create_sheet('Prospects')
    p.append(['Customer Data'])
    p.append(['Date', 'Lead', 'S', 'Phone Number', 'Emai id', 'UR', 'Names', 'DOB', 'Age', 'Enq.', 'onward Dep',
              'Return Date', 'Dep.', 'Arrival', 'TC #1', 'TI #1', 'CL #1', 'LS', 'Next Follow-Up', 'Next Step', 'Comments'])
    p.append([46304, 'Kiran', 'IW', '91 98480 00007', 'kiran@example.com', 'M', 'Lakshmi', None, 67, 'TI', 46340,
              46400, 'HYD', 'Dallas', None, 46305, 'Y', 'R', 46310, 'TI #2', 'Wants comprehensive'])
    p.append([46305, 'Venu', 'FB', '1 (408) 593-6794', None, 'F', None, None, 70, 'TC', 46350, None, 'HYD', 'SFO',
              46305, None, None, 'U', None, None, None])
    for _ in range(5):                                    # the dragged-down formula rows
        p.append([None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None,
                  None, None, 46308, None])
    c = wb.create_sheet('Customer Engagement')
    c.append(['Customer Data'])
    c.append(['Date', 'Customer', 'S', 'Phone Number', 'Emai id', 'UR', 'Names', 'DOB', 'Age', 'Enq.', 'Start Date',
              'End Date', 'TD', 'Dep.', 'Arrival', 'Dep.', 'Arrival', 'Sale', 'PE', 'Claim Number', 'CE #1', 'CE #2'])
    c.append([46265, 'Sree', 'UC', '1.928.981.9840', 'sree@example.com', 'M', None, None, 63, 'TI', 46305, 46400])
    f = wb.create_sheet('Prospect Fallout')
    f.append(['Customer Data'])
    f.append(['Date', 'Lead', 'S', 'Phone Number', 'Emai id', 'Names', 'Enq.', 'onward Dep', 'FM', 'BW', 'F M', 'CM',
              'LS', 'Date', 'Next Step', 'Comments'])
    f.append([46300, 'Gopi', 'I', '+61 434 727 317', None, None, 'TI', None, None, None, None, None, 'DL', None, None,
              'bought elsewhere'])
    wb.create_sheet('KPI').append(['Totals'])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_reading_the_sales_workbook():
    rows, stats, kpi = crm_import.read_sheet(workbook())
    by = {r['name'] or r['details'].get('post'): r for r in rows}
    assert stats['Cold Leads']['rows'] == 2 and stats['Prospects']['rows'] == 2
    assert by['HYD DFW NRI TC']['phone'] == '+18044976837'
    assert by['BLR DTW NRI TC']['identities'] == [('instagram', 'https://www.instagram.com/some.one?igsh=x')]
    kiran = by['Kiran']
    assert kiran['phone'] == '+919848000007' and kiran['product'] == 'TI' and kiran['source'] == 'wa-india'
    assert kiran['temperature'] == 'R' and kiran['next_step'] == 'TI #2' and kiran['departure']
    assert [c for c, _ in kiran['messages']] == ['TI #1', 'CL #1']
    assert by['Venu']['phone'] == '+14085936794' and by['Venu']['product'] == 'TC'
    assert by['Gopi']['phone'] == '+61434727317'


def test_importing_the_sales_workbook(app, client, admin_user, cs_user):
    login(client, admin_user.email)
    r = client.post('/admin/crm/setup/import', data={'sheet': (workbook(), 'Sales Process 1.5.xlsx')},
                    content_type='multipart/form-data')
    page = r.get_data(as_text=True)
    assert 'leads found' in page and 'Prospects' in page
    token = page.split('name="token" value="', 1)[1].split('"', 1)[0]
    client.post('/admin/crm/setup/import/confirm', data={'token': token, 'owner': cs_user.id})
    leads = {l.person.display_name: l for l in CrmLead.query.all()}
    assert len(leads) == 6
    kiran = leads['Kiran']
    assert kiran.stage.key == 'prospect' and kiran.temperature == 'R' and kiran.owner_id == cs_user.id
    assert kiran.open_task.code == 'TI #2'
    assert CrmTouchpoint.query.filter_by(lead_id=kiran.id, direction='out').count() == 2
    assert leads['Venu'].stage.key == 'in_touch'                     # a TC prospect
    assert leads['Sree'].status == 'won' and leads['Sree'].open_task is None   # old customers are not messaged
    assert leads['Gopi'].status == 'lost' and leads['Gopi'].temperature == 'DL'
    assert leads['some.one'].person.identity('instagram') == 'some.one'
    assert Notification.query.filter_by(type='crm_lead').count() == 0
    # the same workbook again adds nothing
    r = client.post('/admin/crm/setup/import', data={'sheet': (workbook(), 'again.xlsx')}, content_type='multipart/form-data')
    token = r.get_data(as_text=True).split('name="token" value="', 1)[1].split('"', 1)[0]
    client.post('/admin/crm/setup/import/confirm', data={'token': token})
    assert CrmLead.query.count() == 6


def test_backfill_button_and_import_tab(app, client, admin_user, cs_user):
    login(client, admin_user.email)
    assert b'Enquiries from before the CRM' in client.get('/admin/crm/setup?tab=import').data
    r = client.post('/admin/crm/setup/backfill', data={'days': 90}, follow_redirects=True)
    assert b'Brought in 0 past enquiries' in r.data
    client.post('/auth/logout')
    login(client, cs_user.email)
    assert client.post('/admin/crm/setup/backfill').status_code == 302 and CrmLead.query.count() == 0


def test_only_the_import_screen_takes_a_big_upload(app):
    with app.test_request_context('/admin/crm/setup/import', method='POST'):
        from flask import request
        assert request.max_content_length == 64 * 1024 * 1024
    with app.test_request_context('/api/contact', method='POST'):
        from flask import request
        assert request.max_content_length == app.config['MAX_CONTENT_LENGTH']


def test_a_known_product_beats_a_catch_all_source(app):
    lead = crm.capture('instagram', name='Neha', identities=[('instagram', '@neha.t')], product='TC', utm={})
    assert lead.pipeline.key == 'tc-companion' and lead.open_task.code == 'TC #1'
    plain = crm.capture('instagram', identities=[('instagram', '@someone.else')], utm={})
    assert plain.pipeline.key == 'general'
