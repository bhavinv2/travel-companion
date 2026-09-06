"""Scraper domain + worker tests. fetchall is replaced by fakes: no Playwright, no browser."""
import io
import json
from datetime import datetime, timedelta

import pytest

from conftest import login, logout
from app import db
from app.models import ScrapeRecipe, ScrapeRun, ScrapeRow, CompanionRequest, ContactPoint, User
from app.services import scraper, scraper_worker, importer
from app.vendor.fetchall.models import Table, Candidate, Inspection
from app.vendor.fetchall.runner import RunStats

COLUMNS = ['title', 'origin', 'flight', 'destination', 'start', 'end', 'role', 'message', 'languages', 'contact',
           '_page', '_url', '_fetched_at', '_key']

RECIPE = {
    'version': 1, 'site': 'example.test', 'start_url': 'https://example.test/list', 'wait_for': 'li.item',
    'list': {'row': 'li.item', 'container': '#list'},
    'fields': {'title': {'sel': 'a.title', 'attr': 'text'}, 'origin': {'sel': 'span.origin', 'attr': 'text'},
               'flight': {'sel': 'span.flight', 'attr': 'text'}, 'destination': {'sel': 'span.dest', 'attr': 'text'}},
    'detail': {'mode': 'popup', 'open': {'sel': 'a.title'}, 'container': '#modal',
               'fields': {'start': {'sel': 'span.start', 'attr': 'text'}, 'end': {'sel': 'span.end', 'attr': 'text'},
                          'role': {'sel': 'span.role', 'attr': 'text'}, 'message': {'sel': 'p.msg', 'attr': 'text'},
                          'languages': {'sel': 'span.lang', 'attr': 'text'}, 'contact': {'sel': 'span.contact', 'attr': 'text'}},
               'close': {'sel': 'span.close', 'scope': 'container'}},
    'paginate': {'mode': 'click', 'click': '#next', 'max_pages': 50, 'stop_after_empty': 3},
}


def fake_rows(n=3, start=0):
    rows = []
    for i in range(start, start + n):
        rows.append({
            'title': f'Companion for mom {i}', 'origin': 'HYD', 'flight': 'QR573', 'destination': 'DFW',
            'start': '12/1/2099, 4:00:00 AM', 'end': 'N/A', 'role': 'Seeking Help' if i % 2 == 0 else 'Offering Help',
            'message': f'Row {i} message', 'languages': 'Telugu, Hindi', 'contact': f'+1 214 555 01{i:02d}',
            '_page': 1, '_url': 'https://example.test/list', '_fetched_at': '2026-08-30T10:00:00+05:30',
            '_key': f'key{i:03d}',
        })
    return rows


class FakeRunner:
    """Stands in for fetchall.runner.run_recipe; honours on_row / should_stop / skip_keys like the real one."""

    def __init__(self, rows):
        self.rows = rows
        self.calls = []

    def __call__(self, recipe, **kw):
        self.calls.append(kw)
        stats = RunStats()
        stats.blank_counts = {c: 0 for c in COLUMNS if not c.startswith('_')}
        out = []
        skip = kw.get('skip_keys') or set()
        for rec in self.rows:
            if kw.get('should_stop') and kw['should_stop']():
                stats.interrupted = True
                stats.stopped_because = 'cancelled'
                break
            if rec['_key'] in skip:
                stats.skipped_known += 1
                continue
            stats.rows_seen += 1
            stats.rows_new += 1
            out.append(dict(rec))
            if kw.get('on_row'):
                kw['on_row'](dict(rec))
            if kw.get('max_rows') and len(out) >= kw['max_rows']:
                stats.stopped_because = f"reached --max-rows {kw['max_rows']}"
                break
        stats.pages = 1
        if kw.get('log'):
            kw['log'](f'  page 1: {len(self.rows)} items, {len(out)} new')
        if not stats.stopped_because:
            stats.stopped_because = 'next control is disabled (last page)'
        return Table(columns=COLUMNS, rows=out), stats


@pytest.fixture()
def scraper_ready(app, monkeypatch):
    app.config.update(SCRAPER_ENABLED=True, SCRAPER_INLINE=True, SCRAPER_FORCE_AVAILABLE=True, SCRAPER_HEADED_TEACH=False)
    runner = FakeRunner(fake_rows(3))
    monkeypatch.setattr(scraper_worker, '_run_recipe', runner)
    return runner


@pytest.fixture()
def recipe(app, admin_user):
    rec = ScrapeRecipe(name='example', site='example.test', start_url=RECIPE['start_url'], mode='recipe',
                       recipe_json=RECIPE, default_source='website', created_by_id=admin_user.id)
    rec.field_mapping = scraper.suggest_mapping(rec.columns)
    db.session.add(rec)
    db.session.commit()
    return rec


# ---------------------------------------------------------------------------
# Mapping & validation
# ---------------------------------------------------------------------------

def test_suggest_mapping_aliases_fuzzy_and_ignore(app):
    m = scraper.suggest_mapping(['title', 'Flying from', 'Departure date', 'contact', '_url', '_key', '_page', 'random_col', 'destinaton'])
    assert m['title'] == 'title' and m['Flying from'] == 'origin' and m['Departure date'] == 'start'
    assert m['contact'] == 'contact' and m['_url'] == 'source_url' and m['_key'] == 'import_key'
    assert m['_page'] == 'ignore' and m['random_col'] == 'ignore'
    assert m['destinaton'] == 'destination'          # fuzzy


def test_apply_mapping_joins_text_and_keeps_meta(app):
    data = {'headline': 'Need help', 'body': 'for my mother', 'from': 'HYD', 'to': 'DFW', 'x': 'ignored',
            '_url': 'https://s/1', '_key': 'k1'}
    mapping = {'headline': 'title', 'body': 'title', 'from': 'origin', 'to': 'destination', 'x': 'ignore'}
    canon = scraper.apply_mapping(data, mapping)
    assert canon['title'] == 'Need help\nfor my mother' and canon['origin'] == 'HYD'
    assert canon['source_url'] == 'https://s/1' and canon['import_key'] == 'k1' and 'x' not in canon


def test_validate_recipe(app):
    assert scraper.validate_recipe(RECIPE) == []
    bad = {'start_url': 'ftp://x', 'list': {}, 'fields': {}, 'detail': {'mode': 'weird'}, 'paginate': {'mode': 'click'}}
    errs = scraper.validate_recipe(bad)
    assert len(errs) == 5 and any('paginate.click' in e for e in errs)
    assert scraper.validate_recipe('nope') == ['The recipe must be a JSON object.']
    assert scraper.recipe_columns(RECIPE) == COLUMNS


def test_availability_states(app, monkeypatch):
    app.config['SCRAPER_FORCE_AVAILABLE'] = False
    assert scraper.availability()['ok'] is False
    app.config['SCRAPER_FORCE_AVAILABLE'] = None
    app.config['SCRAPER_ENABLED'] = False
    assert 'disabled' in scraper.availability()['message']
    app.config['SCRAPER_ENABLED'] = True
    app.config['SCRAPER_FORCE_AVAILABLE'] = True
    assert scraper.availability()['ok'] is True


# ---------------------------------------------------------------------------
# Runs: storage, incremental, cancel, source mode, detect
# ---------------------------------------------------------------------------

def test_run_stores_rows_and_second_run_is_incremental(scraper_ready, recipe, admin_user, db):
    run = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    assert run.status == 'done' and run.rows_new == 3 and run.rows_total == 3
    assert ScrapeRow.query.filter_by(recipe_id=recipe.id).count() == 3
    assert scraper_ready.calls[0]['skip_keys'] == set()
    assert 'page 1' in (run.log or '') and run.progress_pages == 1
    assert run.stats['report']['status'] == 'ok'
    assert recipe.last_status == 'done' and recipe.last_run_at is not None

    scraper_ready.rows = fake_rows(4)                  # one new item appeared on the site
    run2 = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    assert scraper_ready.calls[1]['skip_keys'] == {'key000', 'key001', 'key002'}
    assert run2.rows_new == 1 and ScrapeRow.query.filter_by(recipe_id=recipe.id).count() == 4
    run3 = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user, options={'incremental': False})
    assert scraper_ready.calls[2]['skip_keys'] == set() and run3.rows_new == 0   # unique (recipe, key) held


def test_cancel_keeps_partial_rows(scraper_ready, recipe, admin_user, db, monkeypatch):
    stored = []
    original = scraper_worker._store_row

    def store_and_cancel(run, rec, index, row, columns):
        ok = original(run, rec, index, row, columns)
        stored.append(row['_key'])
        if len(stored) == 2:
            run.cancel_requested = True
            db.session.commit()
        return ok

    monkeypatch.setattr(scraper_worker, '_store_row', store_and_cancel)
    run = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    assert run.status == 'cancelled' and ScrapeRow.query.filter_by(run_id=run.id).count() == 2


def test_rows_matching_existing_posts_are_flagged_duplicate(scraper_ready, recipe, admin_user, db):
    t = CompanionRequest(travel_type='air', trip_type='one_way', flying_from='Hyderabad (HYD)',
                         destination='Dallas (DFW)', import_key='key001', role='seeking_help', source='website')
    t.set_status('unconfirmed')
    db.session.add(t)
    db.session.commit()
    run = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    dup = ScrapeRow.query.filter_by(run_id=run.id, key='key001').first()
    assert dup.status == 'duplicate' and dup.duplicate_of_id == t.id and run.rows_duplicate == 1


def test_single_flight_and_stale_recovery(scraper_ready, recipe, admin_user, db):
    active = ScrapeRun(recipe_id=recipe.id, kind='run', status='running', started_at=datetime.utcnow() - timedelta(hours=1),
                       heartbeat_at=datetime.utcnow() - timedelta(hours=1))
    db.session.add(active)
    db.session.commit()
    with pytest.raises(scraper_worker.AlreadyRunning):
        scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    assert scraper_worker.recover_stale_runs() == 1
    assert db.session.get(ScrapeRun, active.id).status == 'failed'
    run = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)      # now allowed
    assert run.status == 'done'


def test_unavailable_and_headed_teach_guard(app, recipe, admin_user):
    app.config.update(SCRAPER_FORCE_AVAILABLE=False)
    with pytest.raises(scraper.ScraperUnavailable):
        scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    app.config.update(SCRAPER_FORCE_AVAILABLE=True, SCRAPER_HEADED_TEACH=False, SCRAPER_INLINE=True)
    with pytest.raises(scraper.ScraperUnavailable):
        scraper_worker.enqueue('teach', options={'url': 'https://example.test'}, actor=admin_user)


def test_detect_run_collects_ranked_candidates(app, admin_user, monkeypatch):
    app.config.update(SCRAPER_INLINE=True, SCRAPER_FORCE_AVAILABLE=True)
    csv_table = Table(columns=['title', 'origin', 'destination', 'start', 'contact'],
                      rows=[{'title': 'A', 'origin': 'HYD', 'destination': 'DFW', 'start': '2099-01-01', 'contact': 'a@x.com'}] * 3)
    html_table = Table(columns=['c1', 'c2'], rows=[{'c1': '1', 'c2': '2'}] * 2)
    insp = Inspection(url='https://example.test/', final_url='https://example.test/', title='Example',
                      candidates=[Candidate(kind='html-table', source='dom:table[0]', table=html_table),
                                  Candidate(kind='csv', source='https://example.test/data.csv', table=csv_table)],
                      responses_seen=5, visible_text='A HYD DFW 2099-01-01 a@x.com')
    monkeypatch.setattr(scraper_worker, '_inspect', lambda url, **kw: insp)
    run = scraper_worker.enqueue('detect', options={'url': 'https://example.test/'}, actor=admin_user)
    assert run.status == 'done'
    cands = run.result['candidates']
    assert len(cands) == 2 and cands[0]['kind'] == 'csv' and cands[0]['rows'] == 3 and len(cands[0]['sample']) == 3


def test_source_mode_run_uses_fetch_direct(app, admin_user, monkeypatch, db):
    app.config.update(SCRAPER_INLINE=True, SCRAPER_FORCE_AVAILABLE=True)
    table = Table(columns=['title', 'origin', 'destination', 'start', 'contact'],
                  rows=[{'title': f'T{i}', 'origin': 'HYD', 'destination': 'DFW', 'start': '2099-01-01', 'contact': f'a{i}@x.com'} for i in range(4)])
    monkeypatch.setattr(scraper_worker, '_fetch_direct', lambda url: Candidate(kind='csv', source=url, table=table))
    rec = ScrapeRecipe(name='csv-site', site='example.test', start_url='https://example.test/', mode='source',
                       source_json={'source': 'https://example.test/data.csv', 'kind': 'csv', 'page_url': 'https://example.test/',
                                    'columns': table.columns})
    rec.field_mapping = scraper.suggest_mapping(rec.columns)
    db.session.add(rec)
    db.session.commit()
    run = scraper_worker.enqueue('run', recipe=rec, actor=admin_user)
    assert run.status == 'done' and run.rows_new == 4
    row = ScrapeRow.query.filter_by(recipe_id=rec.id).first()
    assert row.key and row.data['_url'] == 'https://example.test/' and row.data['_key'] == row.key
    run2 = scraper_worker.enqueue('run', recipe=rec, actor=admin_user)
    assert run2.rows_new == 0                                   # same content hashes -> no duplicates


def test_enqueue_scheduled_respects_interval(scraper_ready, recipe, admin_user, db):
    recipe.schedule_enabled = True
    recipe.schedule_every_hours = 6
    db.session.commit()
    assert scraper_worker.enqueue_scheduled() == 1
    assert scraper_worker.enqueue_scheduled() == 0              # just ran
    recipe.last_run_at = datetime.utcnow() - timedelta(hours=7)
    db.session.commit()
    assert scraper_worker.enqueue_scheduled() == 1


# ---------------------------------------------------------------------------
# Rows -> posts, export
# ---------------------------------------------------------------------------

def test_create_posts_bulk_is_idempotent(scraper_ready, recipe, admin_user, db):
    run = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    rows = ScrapeRow.query.filter_by(run_id=run.id).order_by(ScrapeRow.id).all()
    preview = scraper.preview_rows(rows, recipe)
    assert all(p['status'] == 'ok' for p in preview)
    assert preview[0]['flying_from'].endswith('(HYD)') and preview[0]['airline'] == 'Qatar Airways'
    assert preview[0]['role'] == 'seeking_help' and preview[1]['role'] == 'offering_help'

    result = scraper.create_posts(recipe, [rows[0].id, rows[1].id], admin_user)
    assert result['created'] == 2 and result['duplicates'] == 0 and result['errors'] == []
    posts = CompanionRequest.query.filter(CompanionRequest.id.in_(result['post_ids'])).all()
    assert all(p.status == 'unconfirmed' and p.source == 'website' and p.source_url == 'https://example.test/list' for p in posts)
    assert {p.import_key for p in posts} == {'key000', 'key001'}
    assert all(not cp.consent_to_share and cp.added_by == 'import' for p in posts for cp in p.contact_points)
    assert all(cp.type == 'mobile' for p in posts for cp in p.contact_points)
    assert 'post_scraped' in [e.event for e in posts[0].events]
    r0 = db.session.get(ScrapeRow, rows[0].id)
    assert r0.status == 'imported' and r0.imported_post_id in result['post_ids'] and r0.imported_by_id == admin_user.id

    again = scraper.create_posts(recipe, [rows[0].id, rows[1].id, rows[2].id], admin_user)
    assert again['created'] == 1 and CompanionRequest.query.count() == 3      # only the third row is new

    # a row whose import_key already exists as a post (e.g. from an Excel import) becomes 'duplicate'
    extra = ScrapeRow(run_id=run.id, recipe_id=recipe.id, row_index=9, key='keyX', data={**fake_rows(1)[0], '_key': 'keyX'}, status='new')
    db.session.add(extra)
    t = CompanionRequest(travel_type='air', trip_type='one_way', flying_from='x', destination='y', import_key='keyX')
    t.set_status('unconfirmed')
    db.session.add(t)
    db.session.commit()
    res = scraper.create_posts(recipe, [extra.id], admin_user)
    assert res['duplicates'] == 1 and db.session.get(ScrapeRow, extra.id).status == 'duplicate'


def test_create_posts_publish_now_and_error_rows(scraper_ready, recipe, admin_user, db):
    scraper_ready.rows = [{**fake_rows(1)[0], 'origin': 'Atlantis', '_key': 'bad1'}, {**fake_rows(1, 5)[0]}]
    run = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    rows = ScrapeRow.query.filter_by(run_id=run.id).order_by(ScrapeRow.id).all()
    res = scraper.create_posts(recipe, [r.id for r in rows], admin_user, publish_now=True)
    assert res['created'] == 1 and len(res['errors']) == 1 and 'Origin' in res['errors'][0]['errors'][0]
    post = db.session.get(CompanionRequest, res['post_ids'][0])
    assert post.status == 'open' and post.claimed_at is not None
    assert db.session.get(ScrapeRow, rows[0].id).status == 'new'      # error rows stay new for fixing
    assert scraper.skip_rows(recipe, [rows[0].id], admin_user) == 1
    assert db.session.get(ScrapeRow, rows[0].id).status == 'skipped'


def test_prefill_form_shapes_request_form(scraper_ready, recipe, admin_user, db):
    run = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    row = ScrapeRow.query.filter_by(run_id=run.id).first()
    form, contact_rows = scraper.prefill_form(row, recipe)
    assert form.get('flying_from').endswith('(HYD)') and form.get('role') == 'seeking_help'
    assert form.getlist('preferred_languages') == ['Telugu', 'Hindi']
    assert contact_rows and contact_rows[0]['type'] == 'mobile' and contact_rows[0]['consent'] is False


def test_export_raw_and_mapped(scraper_ready, recipe, admin_user, db, app):
    import openpyxl
    run = scraper_worker.enqueue('run', recipe=recipe, actor=admin_user)
    rows = ScrapeRow.query.filter_by(run_id=run.id).all()
    raw = scraper.export_workbook(recipe, rows, mapped=False, run=run)
    wb = openpyxl.load_workbook(raw, read_only=True)
    assert wb.sheetnames == ['rows', 'Source']
    header = [c.value for c in next(wb['rows'].iter_rows(min_row=1, max_row=1))]
    assert header[:4] == ['title', 'origin', 'flight', 'destination'] and header[-2:] == ['scrape_status', 'post_id']
    assert wb['rows'].max_row == 4
    mapped = scraper.export_workbook(recipe, rows, mapped=True)
    with open(mapped, 'rb') as f:
        records, err = importer.parse_file(f, 'x.xlsx')
    assert err is None and len(records) == 3
    built = importer.build_row(records[0], 0)
    assert built['status' if 'status' in built else 'errors'] in ('ok', []) or built['errors'] == []
    assert built['flying_from'].endswith('(HYD)') and built['import_key'] == 'key000'


def test_importer_handles_messy_scraped_values(app):
    assert importer.parse_languages('N/A') == []
    assert importer.parse_languages('My Mother Speaks English/ Telugu') == ['English', 'Telugu']
    assert importer.parse_languages('Telugu, Hindi and Kannada') == ['Telugu', 'Hindi', 'Kannada']
    assert importer.parse_flight('KLM', None) == ('KLM Royal Dutch Airlines', None) or importer.parse_flight('KLM', None)[1] is None
    assert importer.parse_flight('Air India', None) == ('Air India', None)
    assert importer.parse_flight('QR 573', None) == ('Qatar Airways', 'QR573')
    assert importer.parse_flight('N/A', 'N/A') == (None, None)
    row = importer.build_row({'origin': 'HYD', 'destination': 'DFW', 'start': '2099-01-01', 'end': 'N/A',
                              'languages': 'n/a', 'contact': '-', 'message': 'NA'}, 0)
    assert row['to_date'] is None and row['languages'] == [] and row['contacts'] == [] and row['message'] == ''


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def test_cli_import_recipe_and_run(app, scraper_ready, admin_user, tmp_path):
    path = tmp_path / 'r.json'
    path.write_text(json.dumps(RECIPE), encoding='utf-8')
    runner = app.test_cli_runner()
    res = runner.invoke(args=['scrape-import-recipe', str(path), '--name', 'clirecipe'])
    assert res.exit_code == 0 and 'clirecipe' in res.output, res.output
    res = runner.invoke(args=['scrape-list'])
    assert 'clirecipe' in res.output
    res = runner.invoke(args=['scrape-run', 'clirecipe', '--max-rows', '2'])
    assert res.exit_code == 0, res.output
    assert 'status=done' in res.output and 'new=2' in res.output
    assert ScrapeRow.query.count() == 2
