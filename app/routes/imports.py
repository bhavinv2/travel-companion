"""CS Excel/CSV import: upload → preview (validated, de-duplicated) → commit as unconfirmed posts."""
from flask import Blueprint, render_template, request, redirect, url_for, flash
from flask_login import login_required, current_user

from app.routes.cs import cs_required, _choices
from app.services import importer

imports_bp = Blueprint('imports', __name__)


@imports_bp.route('/cs/import/template.xlsx')
@login_required
@cs_required
def import_template():
    """A filled-in sample workbook: example rows to copy over + a "How to fill" guide sheet."""
    from io import BytesIO
    from datetime import date, timedelta
    from flask import send_file
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = ['poster_name', 'traveler_name', 'on_behalf_of', 'role', 'origin', 'destination',
               'start', 'end', 'airline', 'flight', 'languages', 'need_help_with',
               'email', 'phone', 'facebook', 'ticket_booked', 'message', 'source_url']
    d1 = (date.today() + timedelta(days=21)).isoformat()
    d2 = (date.today() + timedelta(days=45)).isoformat()
    d3 = (date.today() + timedelta(days=30)).isoformat()
    samples = [
        ['Priya Patel', 'Lakshmi Patel', 'mother', 'seeking_help', 'Hyderabad (HYD)', 'Dallas (DFW)',
         d1, '', 'Qatar Airways', 'QR573', 'Telugu, English', 'wheelchair, first_time',
         'priya@example.com', '+1 555 010 1234', '', 'yes',
         'My mother is travelling alone for the first time. She speaks Telugu and a little English.', ''],
        ['Ravi Kumar', '', 'myself', 'offering_help', 'HYD', 'DFW',
         d1, d2, 'Qatar Airways', 'QR573', 'Telugu, Hindi, English', '',
         '', '+91 98490 12345', '', 'yes',
         'Happy to help a co-passenger with forms and bags at immigration.', ''],
        ['Anita D', '', 'myself', 'open_to_either', 'Mumbai (BOM)', 'London (LHR)',
         d3, '', 'Air India', 'AI131', 'Hindi, Marathi', 'connections',
         'anita@example.com', '', 'facebook.com/anita.d', 'no',
         'Long layover - company would be lovely.', 'https://facebook.com/groups/desitravel/posts/123'],
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = 'Travellers'
    navy = PatternFill('solid', fgColor='0F2340')
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = navy
        c.alignment = Alignment(vertical='center')
    for r in samples:
        ws.append(r)
    for i, w in enumerate([16, 15, 12, 14, 18, 18, 12, 12, 15, 10, 22, 24, 22, 17, 22, 12, 52, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

    hs = wb.create_sheet('How to fill')
    guide = [
        ('Column', 'Required?', 'What to put there'),
        ('poster_name', 'recommended', 'Who posted / who CS talked to.'),
        ('traveler_name', 'optional', 'Only when someone else is travelling (e.g. a parent).'),
        ('on_behalf_of', 'optional', 'myself, mother, father, parents, son, daughter, friend, other'),
        ('role', 'YES', 'seeking_help, offering_help or open_to_either'),
        ('origin / destination', 'YES', 'City, airport name or IATA code - "Hyderabad (HYD)", "HYD" and "Hyderabad" all work.'),
        ('start / end', 'start = YES', 'Departure / return date. Safest format: YYYY-MM-DD (e.g. 2026-09-20). Empty end = one-way.'),
        ('airline / flight', 'optional', 'e.g. Qatar Airways / QR573. The same flight makes the strongest match.'),
        ('languages', 'optional', 'Comma-separated: Telugu, Hindi, English'),
        ('need_help_with', 'optional', 'Comma-separated: wheelchair, toddler, first_time, immigration, connections, language_barrier, medicines, documents'),
        ('email / phone / facebook', 'optional', 'At least one is recommended. Stored UNCONSENTED until the person confirms via the claim link.'),
        ('ticket_booked', 'optional', 'yes / no'),
        ('message', 'optional', 'The free-text story shown on the post.'),
        ('source_url', 'optional', 'Where the post was found - also used for duplicate detection.'),
        ('', '', ''),
        ('Good to know', '', 'Header names are flexible (from/to, date, comments... also work) and extra columns are ignored. '
                             'Every imported post is created UNCONFIRMED - nothing goes public until CS confirms it.'),
    ]
    for row in guide:
        hs.append(row)
    for c in hs[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = navy
    hs.column_dimensions['A'].width = 26
    hs.column_dimensions['B'].width = 14
    hs.column_dimensions['C'].width = 110
    for cells in hs.iter_rows(min_row=2):
        cells[2].alignment = Alignment(wrap_text=True, vertical='top')

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='connecting-desis-import-template.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@imports_bp.route('/cs/import', methods=['GET', 'POST'])
@login_required
@cs_required
def import_page():
    if request.method == 'POST' and request.form.get('step') == 'commit':
        rows = importer.payload_to_rows(request.form.get('payload'))
        selected = set(request.form.getlist('include'))
        for r in rows:
            if str(r.get('idx')) not in selected:
                r['status'] = 'skipped'
        importer.annotate_duplicates([r for r in rows if r.get('status') == 'ok']) if rows else None
        created = importer.commit_rows(rows, current_user)
        flash(f'Imported {len(created)} post(s) as unconfirmed.', 'success')
        return redirect(url_for('cs.posts', status='unconfirmed', sort='newest'))

    if request.method == 'POST':
        f = request.files.get('file')
        if not f or not f.filename:
            flash('Choose an .xlsx or .csv file.', 'danger')
            return redirect(url_for('imports.import_page'))
        default_source = request.form.get('default_source') if request.form.get('default_source') in ('facebook', 'website', 'excel') else 'website'
        records, err = importer.parse_file(f.stream, f.filename)
        if err:
            flash(err, 'danger')
            return redirect(url_for('imports.import_page'))
        rows = importer.annotate_duplicates([importer.build_row(rec, i, default_source) for i, rec in enumerate(records)])
        counts = {s: sum(1 for r in rows if r['status'] == s) for s in ('ok', 'error', 'duplicate')}
        return render_template('cs/import.html', rows=rows, counts=counts, payload=importer.rows_to_payload(rows),
                               filename=f.filename, **_choices())

    return render_template('cs/import.html', rows=None, counts=None, payload=None, filename=None, **_choices())
