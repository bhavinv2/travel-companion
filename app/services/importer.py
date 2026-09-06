"""Excel / CSV import of companion requests (plan §4.4).

The template is a superset of the `fetchall` scraper output
(title, origin, flight, destination, start, end, role, message, languages, contact, _url, _key) so scraper
files upload unchanged; CS can add poster_name, traveler_name, relationship, need_help_with, airline,
pref_gender, pref_age_min/max, email, phone, facebook, source, source_url.

Rows are parsed into plain dicts (JSON-serialisable) so the preview → commit round-trip can carry them in the
form without re-reading the file. Imported posts are always `unconfirmed` and contact points unconsented.
"""
import csv
import hashlib
import io
import json
import re
from datetime import datetime, date

from app import db
from app.models import CompanionRequest, ContactPoint, ActivityEvent, TRIP_ROLES, PREF_GENDERS
from app.services.locations import normalize_location
from app.services.contacts import detect_type, normalize_value, validate

COLUMN_ALIASES = {
    'title': ['title', 'subject'],
    'message': ['message', 'comments', 'additional_comments', 'description', 'notes', 'details'],
    'origin': ['origin', 'from', 'flying_from', 'origin_iata', 'departure_airport'],
    'destination': ['destination', 'to', 'dest', 'destination_iata', 'arrival_airport'],
    'start': ['start', 'from_date', 'departure', 'departure_date', 'date', 'travel_date'],
    'end': ['end', 'to_date', 'return', 'return_date'],
    'flight': ['flight', 'flight_number', 'flight_no', 'flightno'],
    'airline': ['airline', 'carrier'],
    'role': ['role', 'type'],
    'languages': ['languages', 'language', 'preferred_languages'],
    'contact': ['contact', 'contact_value', 'contact_info'],
    'email': ['email', 'e-mail', 'email_address'],
    'phone': ['phone', 'mobile', 'whatsapp', 'phone_number'],
    'facebook': ['facebook', 'fb', 'facebook_profile'],
    'poster_name': ['poster_name', 'name', 'posted_by', 'poster'],
    'traveler_name': ['traveler_name', 'traveller_name', 'traveler', 'traveller'],
    'on_behalf_of': ['on_behalf_of', 'relationship', 'relation'],
    'need_help_with': ['need_help_with', 'traveller_needs', 'traveler_needs', 'needs'],
    'pref_gender': ['pref_gender', 'companion_gender'],
    'pref_age_min': ['pref_age_min', 'age_min'],
    'pref_age_max': ['pref_age_max', 'age_max'],
    'ticket_booked': ['ticket_booked', 'ticket'],
    'source': ['source'],
    'source_url': ['_url', 'source_url', 'url', 'link', 'post_url'],
    'import_key': ['_key', 'import_key', 'key', 'id'],
}
_ALIAS_TO_CANON = {alias: canon for canon, aliases in COLUMN_ALIASES.items() for alias in aliases}

DATE_FORMATS = ('%m/%d/%Y, %I:%M:%S %p', '%m/%d/%Y %I:%M:%S %p', '%m/%d/%Y, %H:%M:%S', '%m/%d/%Y',
                '%Y-%m-%d %H:%M:%S', '%Y-%m-%dT%H:%M:%S', '%Y-%m-%d', '%d-%m-%Y', '%d %b %Y', '%d %B %Y',
                '%b %d, %Y', '%B %d, %Y')
class _LiveLookup:
    """Dict-like view rebuilt from the (admin-configurable) option lists on every access."""

    def __init__(self, build):
        self._build = build

    def get(self, k, d=None):
        return self._build().get(k, d)

    def __getitem__(self, k):
        return self._build()[k]

    def __contains__(self, k):
        return k in self._build()


def _build_lang_lookup():
    from app.options import get_list
    return {l.lower(): l for l in get_list('languages')}


def _build_need_lookup():
    from app.options import get_list
    needs = get_list('traveller_needs')
    d = {label.lower(): key for key, label in needs}
    d.update({key: key for key, _ in needs})
    return d


_LANG_LOOKUP = _LiveLookup(_build_lang_lookup)
_NEED_LOOKUP = _LiveLookup(_build_need_lookup)

def _airline_name(code):
    from app.services import airlines
    return airlines.name_for(code)


def _norm_header(h):
    return re.sub(r'[^a-z0-9_]+', '_', str(h or '').strip().lower()).strip('_')


def _cell(v):
    if v is None:
        return ''
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    s = str(v).strip()
    return '' if s.lower() in ('n/a', 'na', 'none', 'null', '-') else s


def parse_file(fileobj, filename):
    """Return (rows, error) where rows is a list of {canonical_column: value}."""
    name = (filename or '').lower()
    raw = fileobj.read()
    if name.endswith('.xlsx') or name.endswith('.xlsm'):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as e:
            return [], f'Could not read the Excel file: {e}'
        ws = wb['rows'] if 'rows' in wb.sheetnames else wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        try:
            headers = [_norm_header(h) for h in next(it)]
        except StopIteration:
            return [], 'The sheet is empty.'
        records = [dict(zip(headers, [_cell(v) for v in r])) for r in it if any(v not in (None, '') for v in r)]
    elif name.endswith('.csv') or name.endswith('.txt'):
        text = raw.decode('utf-8-sig', errors='replace')
        reader = csv.DictReader(io.StringIO(text))
        records = [{_norm_header(k): _cell(v) for k, v in r.items() if k is not None} for r in reader]
    else:
        return [], 'Upload an .xlsx or .csv file.'
    rows = []
    for rec in records:
        canon = {}
        for k, v in rec.items():
            c = _ALIAS_TO_CANON.get(k)
            if c and v and c not in canon:
                canon[c] = v
        rows.append(canon)
    return rows, None


def parse_date(s):
    if not s:
        return None
    s = str(s).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s).date()
    except ValueError:
        return None


def parse_role(s):
    s = (s or '').strip().lower()
    if not s:
        return 'seeking_help'
    if 'offer' in s or 'help' in s and 'can' in s or s == 'offering_help':
        return 'offering_help'
    if 'seek' in s or 'need' in s or 'look' in s or s == 'seeking_help':
        return 'seeking_help'
    if s in TRIP_ROLES:
        return s
    return 'open'


_EMPTY_VALUES = {'n/a', 'na', 'none', 'null', '-', '--', 'nil', 'not available'}


def is_empty_value(s):
    return not s or str(s).strip().lower() in _EMPTY_VALUES


def parse_languages(s):
    if is_empty_value(s):
        return []
    out = []
    for part in re.split(r'[,/&;|]|\band\b', s or ''):
        p = part.strip().strip('.')
        if len(p) < 2 or p.lower() in _EMPTY_VALUES:
            continue
        known = _LANG_LOOKUP.get(p.lower())
        if known:
            out.append(known)
            continue
        # free text like "My mother speaks Telugu" -> keep any known language words inside it
        found = [_LANG_LOOKUP[w] for w in re.findall(r'[A-Za-z]+', p.lower()) if w in _LANG_LOOKUP]
        out.extend(found or ([p.title()] if len(p.split()) <= 2 else []))
    return list(dict.fromkeys(out))


def parse_needs(s):
    out = []
    for part in re.split(r'[,/;|]', s or ''):
        p = part.strip().lower()
        if p and p in _NEED_LOOKUP:
            out.append(_NEED_LOOKUP[p])
    return out


def parse_flight(flight, airline):
    """'QR573' -> ('Qatar Airways', 'QR573'); explicit airline column wins; a value with no digits
    (e.g. 'KLM', 'Air India') is an airline name, not a flight number."""
    flight = '' if is_empty_value(flight) else (flight or '').strip()
    airline = '' if is_empty_value(airline) else (airline or '').strip()
    if flight and not re.search(r'\d', flight):
        airline = airline or (_airline_name(flight) if len(flight) <= 3 else None) or flight
        flight = ''
    flight = flight.upper().replace(' ', '')
    if not airline and flight:
        m = re.match(r'^([A-Z0-9]{2})\s*\d{1,4}[A-Z]?$', flight)
        if m:
            airline = _airline_name(m.group(1)) or ''
    return airline or None, flight or None


def build_row(canon, idx, default_source='website'):
    """Turn one parsed record into a normalised, validated import row."""
    canon = {k: ('' if isinstance(v, str) and is_empty_value(v) else v) for k, v in (canon or {}).items()}
    row = {'idx': idx, 'errors': [], 'warnings': [], 'contacts': []}
    origin = normalize_location(canon.get('origin', ''))
    dest = normalize_location(canon.get('destination', ''))
    from_date = parse_date(canon.get('start'))
    to_date = parse_date(canon.get('end'))
    airline, flight_number = parse_flight(canon.get('flight'), canon.get('airline'))

    row.update({
        'title': canon.get('title', '')[:300],
        'message': canon.get('message', ''),
        'origin_text': canon.get('origin', ''), 'destination_text': canon.get('destination', ''),
        'flying_from': origin['display'] if origin else canon.get('origin', ''),
        'destination': dest['display'] if dest else canon.get('destination', ''),
        'from_date': from_date.isoformat() if from_date else None,
        'to_date': to_date.isoformat() if to_date and from_date and to_date > from_date else None,
        'airline': airline, 'flight_number': flight_number,
        'role': parse_role(canon.get('role')),
        'languages': parse_languages(canon.get('languages')),
        'needs': parse_needs(canon.get('need_help_with')),
        'poster_name': canon.get('poster_name', '')[:120],
        'traveler_name': canon.get('traveler_name', '')[:120],
        'on_behalf_of': canon.get('on_behalf_of', '').lower().replace(' ', '_')[:50],
        'pref_gender': canon.get('pref_gender', '').lower() if canon.get('pref_gender', '').lower() in PREF_GENDERS else 'any',
        'pref_age_min': _int(canon.get('pref_age_min')), 'pref_age_max': _int(canon.get('pref_age_max')),
        'ticket_booked': (canon.get('ticket_booked', '') or '').lower() in ('yes', 'y', 'true', '1', 'booked'),
        'source': canon.get('source', '').lower() if canon.get('source', '').lower() in ('facebook', 'website', 'excel') else default_source,
        'source_url': canon.get('source_url', '')[:500],
    })
    for col in ('contact', 'email', 'phone', 'facebook'):
        v = canon.get(col, '')
        if not v:
            continue
        ctype = {'email': 'email', 'phone': 'mobile', 'facebook': 'facebook'}.get(col) or detect_type(v)
        ok, err = validate(ctype, v)
        if ok:
            row['contacts'].append({'type': ctype, 'value': normalize_value(ctype, v)})
        else:
            row['warnings'].append(f'{col}: {err}')
    key = canon.get('import_key') or ''
    if not key:
        basis = '|'.join([row['source_url'], row['flying_from'], row['destination'], row['from_date'] or '',
                          ';'.join(c['value'] for c in row['contacts'])])
        key = hashlib.sha1(basis.encode('utf-8')).hexdigest()
    row['import_key'] = key[:64]

    if not origin:
        row['errors'].append(f"Origin not recognised: {canon.get('origin') or '(empty)'}")
    if not dest:
        row['errors'].append(f"Destination not recognised: {canon.get('destination') or '(empty)'}")
    if not from_date:
        row['errors'].append(f"Departure date not understood: {canon.get('start') or '(empty)'}")
    elif from_date < date.today():
        row['warnings'].append('Departure date is in the past')
    if not row['contacts']:
        row['warnings'].append('No contact information — the person cannot be reached')
    return row


def _int(v):
    try:
        return int(str(v).strip()) if v not in (None, '') else None
    except ValueError:
        return None


def annotate_duplicates(rows):
    """Mark rows already imported (same import_key) or whose contact value exists on another post."""
    keys = [r['import_key'] for r in rows]
    existing = {t.import_key: t.id for t in CompanionRequest.query.filter(CompanionRequest.import_key.in_(keys)).all()} if keys else {}
    values = [c['value'] for r in rows for c in r['contacts']]
    known = {}
    if values:
        for cp in ContactPoint.query.filter(ContactPoint.value.in_(values), ContactPoint.trip_id.isnot(None)).all():
            known.setdefault(cp.value, cp.trip_id)
    seen_keys = set()
    for r in rows:
        if r['import_key'] in existing:
            r['duplicate_of'] = existing[r['import_key']]
            r['status'] = 'duplicate'
        elif r['import_key'] in seen_keys:
            r['status'] = 'duplicate'
            r['warnings'].append('Same row appears twice in the file')
        else:
            r['status'] = 'error' if r['errors'] else 'ok'
        seen_keys.add(r['import_key'])
        for c in r['contacts']:
            if c['value'] in known:
                r['warnings'].append(f"Contact {c['value']} already on post #{known[c['value']]}")
    return rows


def commit_rows(rows, actor):
    """Create unconfirmed posts for rows with status 'ok'. Returns list of created posts."""
    from app.services.locations import apply_route
    from app.services import matching
    created = []
    for r in rows:
        if r.get('status') != 'ok':
            continue
        if CompanionRequest.query.filter_by(import_key=r['import_key']).first():
            continue
        comments = '\n'.join(x for x in (r.get('title'), r.get('message')) if x).strip() or None
        t = CompanionRequest(
            user_id=None, created_by_id=actor.id, source=r['source'] or 'excel', source_url=r['source_url'] or None,
            import_key=r['import_key'], travel_type='air', trip_type='round_trip' if r['to_date'] else 'one_way',
            poster_name=r['poster_name'] or None, traveler_name=r['traveler_name'] or None,
            on_behalf_of=r['on_behalf_of'] or None, role=r['role'],
            flying_from=r['flying_from'], destination=r['destination'],
            from_date=date.fromisoformat(r['from_date']), to_date=date.fromisoformat(r['to_date']) if r['to_date'] else None,
            airline=r['airline'], flight_number=r['flight_number'],
            preferred_languages=r['languages'], traveller_needs=r['needs'],
            pref_gender=r['pref_gender'], pref_age_min=r['pref_age_min'], pref_age_max=r['pref_age_max'],
            ticket_booked=r['ticket_booked'], additional_comments=comments,
        )
        t.expires_at = t.to_date or t.from_date
        t.set_status('unconfirmed')
        apply_route(t)
        db.session.add(t)
        db.session.flush()
        for c in r['contacts']:
            db.session.add(ContactPoint(trip=t, type=c['type'], value=c['value'], consent_to_share=False, added_by='import'))
        ActivityEvent.log('post_imported', t, actor=actor, source=t.source, import_key=t.import_key)
        created.append(t)
    db.session.commit()
    for t in created:
        matching.compute_matches_for(t, include_unconfirmed=True, actor=actor)
    return created


def rows_to_payload(rows):
    return json.dumps(rows, separators=(',', ':'))


def payload_to_rows(payload):
    try:
        rows = json.loads(payload or '[]')
    except ValueError:
        return []
    return rows if isinstance(rows, list) else []
